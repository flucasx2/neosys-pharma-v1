"""Lecture stricte de listes : aucune adhésion ou vérification de droits déduite."""
import base64
import csv
import hashlib
import io
import zipfile
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException
from datetime import date,datetime
from uuid import UUID

COLONNES=('reference_externe','nom','date_naissance','client_id','beneficiaire_id','revision_beneficiaire')
MAX_LIGNES=5000


def lire_liste(contenu_base64,format,feuille=None):
    try:raw=base64.b64decode(contenu_base64,validate=True)
    except (ValueError,TypeError) as exc:raise ValueError('Fichier encodé invalide.') from exc
    if not raw or len(raw)>4_000_000:raise ValueError('Fichier vide ou supérieur à 4 Mo.')
    if format=='CSV':
        try:text=raw.decode('utf-8-sig')
        except UnicodeDecodeError as exc:raise ValueError('CSV UTF-8 requis.') from exc
        separator=';' if ';' in text.partition('\n')[0] else ','
        values=csv.reader(io.StringIO(text,newline=''),delimiter=separator,strict=True)
    elif format=='XLSX':
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                entries=archive.infolist()
                if len(entries)>2000 or sum(x.file_size for x in entries)>32_000_000 or any(x.flag_bits&1 for x in entries):raise ValueError('Classeur trop volumineux ou chiffré.')
            from openpyxl import load_workbook
            workbook=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
        except (zipfile.BadZipFile,KeyError,OSError,ParseError,DefusedXmlException) as exc:raise ValueError('Classeur XLSX invalide.') from exc
        def excel():
            try:
                if feuille:
                    if feuille not in workbook.sheetnames:raise ValueError('Feuille Excel absente.')
                    sheet=workbook[feuille]
                else:
                    if len(workbook.sheetnames)!=1:raise ValueError('Choisir explicitement une feuille pour ce classeur multiple.')
                    sheet=workbook.active
                if (sheet.max_column or 0)>len(COLONNES):raise ValueError('Colonnes Excel inconnues : utiliser le modèle fourni.')
                sheet.reset_dimensions()
                for row in sheet.iter_rows(max_col=len(COLONNES)):
                    if any(cell.data_type=='f' for cell in row):raise ValueError('Les formules Excel sont interdites ; fournir des valeurs vérifiées.')
                    yield [cell.value for cell in row]
            finally:workbook.close()
        values=excel()
    else:raise ValueError('Format CSV ou XLSX requis.')
    try:
        headers=next(values,None)
        if headers is None:raise ValueError('En-têtes obligatoires.')
        while headers and headers[-1] is None:headers.pop()
        if len(headers)!=len(set(headers)) or not {'reference_externe','nom'}<=set(headers) or set(headers)-set(COLONNES):
            raise ValueError('En-têtes uniques reference_externe, nom et colonnes facultatives documentées requis. Aucune colonne de droits acceptée.')
        result=[];seen=set()
        for number,cells in enumerate(values,2):
            if all(c is None or c=='' for c in cells):continue
            if len(result)>=MAX_LIGNES:raise ValueError('Au plus 5 000 lignes par fichier ; scinder les listes plus grandes.')
            errors=[]
            if len(cells)>len(headers) and any(c is not None and c!='' for c in cells[len(headers):]):errors.append('Champs sans en-tête.')
            if len(cells)<len(headers):errors.append('Nombre de champs incorrect.')
            row=dict(zip(headers,cells));data={}
            for key in ('reference_externe','nom'):
                value=row.get(key)
                if not isinstance(value,str) or not value.strip() or len(value)>200 or any(ord(c)<32 for c in value):errors.append(key+' : texte non vide requis (identifiants Excel en texte pour conserver les zéros).')
                data[key]=value.strip() if isinstance(value,str) else ''
            for key in ('client_id','beneficiaire_id'):
                value=row.get(key)
                try:data[key]=str(UUID(str(value))) if value not in ('',None) else None
                except (ValueError,TypeError):errors.append(key+' : UUID invalide.');data[key]=None
            value=row.get('date_naissance')
            try:data['date_naissance']=(value.date() if isinstance(value,datetime) else value if isinstance(value,date) else date.fromisoformat(value)).isoformat() if value not in ('',None) else None
            except (ValueError,TypeError):errors.append('date_naissance : AAAA-MM-JJ ou date Excel requise.');data['date_naissance']=None
            rev=row.get('revision_beneficiaire')
            if data['beneficiaire_id']:
                try:
                    if isinstance(rev,bool) or not str(rev).isdigit() or int(rev)<1:raise ValueError()
                    data['revision_beneficiaire']=int(rev)
                except (ValueError,TypeError):errors.append('Révision positive du bénéficiaire existant requise.');data['revision_beneficiaire']=None
            else:
                data['revision_beneficiaire']=None
                if rev not in (None,''):errors.append('Révision sans bénéficiaire existant.')
            if data['reference_externe'] in seen:errors.append('Référence dupliquée dans le fichier.')
            seen.add(data['reference_externe'])
            result.append({'numero':number,'donnees':data,'erreurs':errors})
        if not result:raise ValueError('Le fichier ne contient aucune personne.')
        return hashlib.sha256(raw).hexdigest(),result
    except (csv.Error,ParseError,DefusedXmlException,zipfile.BadZipFile) as exc:raise ValueError('Structure de liste invalide.') from exc
    finally:
        close=getattr(values,'close',None)
        if close:close()
