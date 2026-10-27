"""Import de personnes, provenance et atomicité ; aucune ouverture de droits."""
import base64
from io import BytesIO
from uuid import UUID,uuid4
import pytest
from sqlalchemy import select,func,text
from sqlalchemy.exc import DBAPIError
from openpyxl import Workbook
from test_stock import raw_client,account,db
from test_p5_conf_conventions import organism
from test_p5a_tiers_payants import root,post
from test_p4p5_groupe_organismes import cible_isolee
from neosys_pharma.domain.import_beneficiaires import lire_liste
from neosys_pharma.infrastructure.tiers_payants_models import BeneficiaireCouvertureModel,AdhesionOrganismeModel,OrigineImportBeneficiaireModel
pytest_plugins=('test_vente',)

def body(payer,content):return {'organisme_id':payer['id'],'nom_fichier':'liste-fictive.csv','format':'CSV','contenu_base64':base64.b64encode(content.encode()).decode()}
def request(c,a,path,data,status=200,key=None):
    r=c.post(root(a)+'/imports/beneficiaires/'+path,json={'cle':str(key or uuid4()),**data});assert r.status_code==status,r.text;return r.json()

def test_csv_previsualisation_confirmation_reimport_et_origines(raw_client,account):
    c,a=raw_client,account;payer=organism(c,a)
    data=body(payer,'reference_externe;nom;date_naissance\n000007;Personne fictive;2000-01-02\n000008;Autre personne;2001-03-04\n')
    preview=request(c,a,'previsualiser',data);assert preview['confirmation_possible'] and preview['compteurs']['NOUVEAU']==2
    assert c.get(root(a)+'/beneficiaires').json()['total']==0
    key=uuid4();payload=data|{'empreinte_attendue':preview['empreinte'],'confirmer':True}
    result=request(c,a,'confirmer',payload,key=key);assert len(result['crees'])==2 and result['droits_confirmes'] is False
    assert request(c,a,'confirmer',payload,key=key)==result
    repeat=request(c,a,'confirmer',payload);assert len(repeat['deja_importes'])==2 and not repeat['crees']
    changed=body(payer,'reference_externe;nom\n000007;Nom changé\n')
    blocked=request(c,a,'previsualiser',changed);assert blocked['compteurs']['INVALIDE']==1
    request(c,a,'confirmer',changed|{'confirmer':True,'empreinte_attendue':blocked['empreinte']},422)
    assert c.get(root(a)+'/adhesions').json()['total']==0
    gen=db(c);s=next(gen)
    try:
        assert s.scalar(select(func.count()).select_from(BeneficiaireCouvertureModel))==2
        assert s.scalar(select(func.count()).select_from(OrigineImportBeneficiaireModel))==2
        assert {row.reference_externe for row in s.scalars(select(BeneficiaireCouvertureModel))}=={'000007','000008'}
        with pytest.raises(DBAPIError):s.execute(text('DELETE FROM tiers_payant_origines_import_beneficiaires'));s.commit()
        s.rollback()
    finally:gen.close()


def test_blocage_global_doublons_hash_revision_scope_et_droits_interdits(raw_client,account):
    c,a=raw_client,account;payer=organism(c,a)
    bad=body(payer,'reference_externe;nom\n001;Correct fictif\n001;Doublon\n')
    preview=request(c,a,'previsualiser',bad);assert not preview['confirmation_possible']
    request(c,a,'confirmer',bad|{'confirmer':True,'empreinte_attendue':preview['empreinte']},422)
    good=body(payer,'reference_externe;nom\n009;Correct fictif\n');preview=request(c,a,'previsualiser',good)
    request(c,a,'confirmer',good|{'confirmer':True,'empreinte_attendue':'0'*64},409)
    request(c,a,'confirmer',good|{'confirmer':False,'empreinte_attendue':preview['empreinte']},422)
    assert c.get(root(a)+'/beneficiaires').json()['total']==0
    unknown=body(payer,'reference_externe;nom;droits_statut\n009;Fictif;CONFIRMES\n');request(c,a,'previsualiser',unknown,422)
    response=c.post(root(a)+'/imports/beneficiaires/previsualiser',json={'cle':str(uuid4()),**good},headers={'X-Site-ID':str(a['other'])});assert response.status_code==403
    existing=post(c,a,'/beneficiaires',{'nom':'Bénéficiaire existant'})
    data=body(payer,f'reference_externe;nom;beneficiaire_id;revision_beneficiaire\nREF-P;Bénéficiaire existant;{existing["id"]};1\n')
    preview=request(c,a,'previsualiser',data);assert preview['compteurs']['RATTACHEMENT']==1
    done=request(c,a,'confirmer',data|{'confirmer':True,'empreinte_attendue':preview['empreinte']});assert done['rattaches']==[existing['id']] and not done['crees']


def test_xlsx_types_formules_dates_et_zeros():
    wb=Workbook();ws=wb.active;ws.append(['reference_externe','nom','date_naissance']);ws.append(['000007','Fictif','2000-01-02'])
    out=BytesIO();wb.save(out)
    digest,rows=lire_liste(base64.b64encode(out.getvalue()).decode(),'XLSX')
    assert len(digest)==64 and rows[0]['donnees']['reference_externe']=='000007' and not rows[0]['erreurs']
    ws['A2']=7;out=BytesIO();wb.save(out)
    _,rows=lire_liste(base64.b64encode(out.getvalue()).decode(),'XLSX');assert rows[0]['erreurs']
    ws['A2']='=1+1';out=BytesIO();wb.save(out)
    with pytest.raises(ValueError,match='formules'):lire_liste(base64.b64encode(out.getvalue()).decode(),'XLSX')
    wb.create_sheet('Autre');out=BytesIO();wb.save(out)
    with pytest.raises(ValueError,match='feuille'):lire_liste(base64.b64encode(out.getvalue()).decode(),'XLSX')


def test_rollback_import_et_reference_existante_sans_fusion(raw_client,account,monkeypatch):
    from neosys_pharma.application.tiers_payants_service import TiersPayantsService
    c,a=raw_client,account;payer=organism(c,a)
    post(c,a,'/beneficiaires',{'nom':'Déjà saisi','reference_externe':'OLD'})
    data=body(payer,'reference_externe;nom\nOLD;Autre nom fictif\n')
    assert request(c,a,'previsualiser',data)['compteurs']['INVALIDE']==1
    data=body(payer,'reference_externe;nom\nQA-01;Premier\nQA-02;Second\n');preview=request(c,a,'previsualiser',data)
    original=TiersPayantsService.creer_beneficiaire;count=0
    def fail(self,payload):
        nonlocal count
        count+=1
        if count==2:raise RuntimeError('Panne import provoquée')
        return original(self,payload)
    monkeypatch.setattr(TiersPayantsService,'creer_beneficiaire',fail)
    with pytest.raises(RuntimeError,match='provoquée'):request(c,a,'confirmer',data|{'confirmer':True,'empreinte_attendue':preview['empreinte']})
    assert c.get(root(a)+'/beneficiaires').json()['total']==1
    gen=db(c);s=next(gen)
    try:assert s.scalar(select(func.count()).select_from(OrigineImportBeneficiaireModel))==0
    finally:gen.close()


def test_permissions_import_et_organisme_hors_organisation(raw_client,account):
    from neosys_pharma.infrastructure.securite_models import RolePermissionModel
    from neosys_pharma.domain.organisation.permissions_systeme import permission_id
    c,a=raw_client,account;payer=organism(c,a);data=body(payer,'reference_externe;nom\nNEW;Fictif\n')
    request(c,a,'previsualiser',data|{'organisme_id':str(uuid4())},404)
    gen=db(c);s=next(gen)
    try:s.execute(RolePermissionModel.__table__.delete().where(RolePermissionModel.permission_id==permission_id('tiers_payant.adhesion.gerer')));s.commit()
    finally:gen.close()
    request(c,a,'previsualiser',data,403)
